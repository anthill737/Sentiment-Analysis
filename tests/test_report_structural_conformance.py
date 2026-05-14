"""P5-T4: Report structural conformance tests.

AC5 — GET /api/jobs/{id}/download/pdf for the mocked-completed Job returns
      a non-empty, openable PDF containing:
        (1) a verdict label and numeric score on or near the first page,
        (2) an executive-summary section with pursue-reasons and pass-reasons text,
        (3) at least one embedded chart image,
        (4) at least one evidence-themed section with body text under a heading.

AC6 — GET /api/jobs/{id}/download/docx for the same Job returns a non-empty,
      openable Word document containing the same four structural elements.
"""

import io
import shutil
from datetime import datetime
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

_FIXTURES_DIR = Path(__file__).parent / "fixtures" / "skill_outputs"

# Expected content from the fixture executive.json / scores.json
_VERDICT = "Worth Exploring"
_SCORE = "7.9"
_PURSUE_SNIPPET = "Strong demand from underserved SMB RIA segment"
_PASS_SNIPPET = "Established players could add SMB tier quickly"
_EVIDENCE_HEADINGS = ["Demand Validation", "Competitive Landscape"]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _insert_done_job(job_id: str, run_dir: Path) -> None:
    from app.database import create_db_and_tables, engine
    from app.models import Job
    from sqlmodel import Session

    create_db_and_tables()
    with Session(engine) as session:
        existing = session.get(Job, job_id)
        if existing:
            session.delete(existing)
            session.commit()
        session.add(
            Job(
                id=job_id,
                subject="structural conformance test",
                payload_json="{}",
                status="done",
                run_dir=str(run_dir),
                verdict=_VERDICT,
                score=7.9,
                completed_at=datetime.utcnow(),
            )
        )
        session.commit()


def _clear_job(job_id: str) -> None:
    from app.database import engine
    from app.models import Job
    from sqlmodel import Session

    with Session(engine) as session:
        job = session.get(Job, job_id)
        if job:
            session.delete(job)
            session.commit()


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture
def auth_client():
    from app.main import app

    with TestClient(app) as c:
        c.post("/auth/login", json={"password": "test-password"})
        yield c


@pytest.fixture
def done_job(tmp_path):
    """Done job in DB backed by real fixture report files copied to a temp run dir."""
    job_id = "p5t4-struct-conf"
    run_dir = tmp_path / "runs" / job_id
    run_dir.mkdir(parents=True)

    for fname in ("report.pdf", "report.docx", "executive.json", "scores.json"):
        src = _FIXTURES_DIR / fname
        if src.exists():
            shutil.copy2(str(src), str(run_dir / fname))

    _insert_done_job(job_id, run_dir)
    yield job_id
    _clear_job(job_id)


# ---------------------------------------------------------------------------
# AC5 — PDF structural conformance
# ---------------------------------------------------------------------------


def test_pdf_download_is_non_empty_and_openable(done_job, auth_client):
    """AC5 base: PDF endpoint returns non-empty bytes that open as a valid PDF."""
    import fitz

    resp = auth_client.get(f"/api/jobs/{done_job}/download/pdf")
    assert resp.status_code == 200, resp.text
    body = resp.content
    assert len(body) > 0, "PDF body is empty"

    doc = fitz.open(stream=body, filetype="pdf")
    assert doc.page_count >= 1, "PDF has no pages"
    doc.close()


def test_pdf_first_page_contains_verdict_and_score(done_job, auth_client):
    """AC5(1): PDF first page contains the verdict label and numeric score."""
    import fitz

    resp = auth_client.get(f"/api/jobs/{done_job}/download/pdf")
    assert resp.status_code == 200
    doc = fitz.open(stream=resp.content, filetype="pdf")
    page0_text = doc.get_page_text(0)
    doc.close()

    assert _VERDICT in page0_text, (
        f"Verdict '{_VERDICT}' not found on PDF page 1. Got: {page0_text[:400]!r}"
    )
    assert _SCORE in page0_text, (
        f"Score '{_SCORE}' not found on PDF page 1. Got: {page0_text[:400]!r}"
    )


def test_pdf_contains_executive_summary_pursue_and_pass_reasons(done_job, auth_client):
    """AC5(2): PDF contains executive-summary section with pursue and pass reason text."""
    import fitz

    resp = auth_client.get(f"/api/jobs/{done_job}/download/pdf")
    assert resp.status_code == 200
    doc = fitz.open(stream=resp.content, filetype="pdf")
    all_text = "".join(doc.get_page_text(i) for i in range(doc.page_count))
    doc.close()

    assert _PURSUE_SNIPPET in all_text, (
        f"Pursue reason not found in PDF. Snippet: {_PURSUE_SNIPPET!r}"
    )
    assert _PASS_SNIPPET in all_text, (
        f"Pass reason not found in PDF. Snippet: {_PASS_SNIPPET!r}"
    )


def test_pdf_contains_at_least_one_embedded_chart_image(done_job, auth_client):
    """AC5(3): PDF contains at least one embedded chart image."""
    import fitz

    resp = auth_client.get(f"/api/jobs/{done_job}/download/pdf")
    assert resp.status_code == 200
    doc = fitz.open(stream=resp.content, filetype="pdf")
    total_images = sum(len(doc.get_page_images(i)) for i in range(doc.page_count))
    doc.close()

    assert total_images >= 1, f"PDF has no embedded images (found {total_images})"


def test_pdf_contains_evidence_section_with_body_text(done_job, auth_client):
    """AC5(4): PDF contains at least one evidence-themed section heading with body text."""
    import fitz

    resp = auth_client.get(f"/api/jobs/{done_job}/download/pdf")
    assert resp.status_code == 200
    doc = fitz.open(stream=resp.content, filetype="pdf")
    all_text = "".join(doc.get_page_text(i) for i in range(doc.page_count))
    doc.close()

    found = any(h in all_text for h in _EVIDENCE_HEADINGS)
    assert found, (
        f"No evidence-themed heading found in PDF. "
        f"Expected one of {_EVIDENCE_HEADINGS}. Text excerpt: {all_text[:500]!r}"
    )


# ---------------------------------------------------------------------------
# AC6 — DOCX structural conformance
# ---------------------------------------------------------------------------


def test_docx_download_is_non_empty_and_openable(done_job, auth_client):
    """AC6 base: DOCX endpoint returns non-empty bytes that open as a valid Word doc."""
    import docx as docxlib

    resp = auth_client.get(f"/api/jobs/{done_job}/download/docx")
    assert resp.status_code == 200, resp.text
    body = resp.content
    assert len(body) > 0, "DOCX body is empty"

    doc = docxlib.Document(io.BytesIO(body))
    assert len(doc.paragraphs) >= 1, "DOCX has no paragraphs"


def test_docx_contains_verdict_and_score(done_job, auth_client):
    """AC6(1): DOCX contains the verdict label and numeric score."""
    import docx as docxlib

    resp = auth_client.get(f"/api/jobs/{done_job}/download/docx")
    assert resp.status_code == 200
    doc = docxlib.Document(io.BytesIO(resp.content))
    all_text = "\n".join(p.text for p in doc.paragraphs)

    assert _VERDICT in all_text, (
        f"Verdict '{_VERDICT}' not found in DOCX. Got: {all_text[:400]!r}"
    )
    assert _SCORE in all_text, (
        f"Score '{_SCORE}' not found in DOCX. Got: {all_text[:400]!r}"
    )


def test_docx_contains_executive_summary_pursue_and_pass_reasons(done_job, auth_client):
    """AC6(2): DOCX contains executive-summary section with pursue and pass reason text."""
    import docx as docxlib

    resp = auth_client.get(f"/api/jobs/{done_job}/download/docx")
    assert resp.status_code == 200
    doc = docxlib.Document(io.BytesIO(resp.content))
    all_text = "\n".join(p.text for p in doc.paragraphs)

    assert _PURSUE_SNIPPET in all_text, (
        f"Pursue reason not found in DOCX. Snippet: {_PURSUE_SNIPPET!r}"
    )
    assert _PASS_SNIPPET in all_text, (
        f"Pass reason not found in DOCX. Snippet: {_PASS_SNIPPET!r}"
    )


def test_docx_contains_at_least_one_embedded_chart_image(done_job, auth_client):
    """AC6(3): DOCX contains at least one embedded chart image."""
    import docx as docxlib

    resp = auth_client.get(f"/api/jobs/{done_job}/download/docx")
    assert resp.status_code == 200
    doc = docxlib.Document(io.BytesIO(resp.content))
    num_images = len(doc.inline_shapes)

    assert num_images >= 1, (
        f"DOCX has no embedded images (InlineShapes count: {num_images})"
    )


def test_docx_contains_evidence_section_with_body_text(done_job, auth_client):
    """AC6(4): DOCX contains an evidence-themed heading followed by body-text paragraphs."""
    import docx as docxlib

    resp = auth_client.get(f"/api/jobs/{done_job}/download/docx")
    assert resp.status_code == 200
    doc = docxlib.Document(io.BytesIO(resp.content))
    paragraphs = doc.paragraphs

    all_text = "\n".join(p.text for p in paragraphs)
    has_heading = any(h in all_text for h in _EVIDENCE_HEADINGS)
    assert has_heading, (
        f"No evidence-themed heading found in DOCX. "
        f"Expected one of {_EVIDENCE_HEADINGS}. Text: {all_text[:500]!r}"
    )

    found_body_after_heading = False
    for i, para in enumerate(paragraphs):
        if para.style.name.startswith("Heading") and any(
            h in para.text for h in _EVIDENCE_HEADINGS
        ):
            for j in range(i + 1, min(i + 5, len(paragraphs))):
                following = paragraphs[j]
                if following.text.strip() and not following.style.name.startswith(
                    "Heading"
                ):
                    found_body_after_heading = True
                    break
            if found_body_after_heading:
                break

    assert found_body_after_heading, (
        "No non-empty body paragraph found immediately after an evidence-themed heading in DOCX"
    )
